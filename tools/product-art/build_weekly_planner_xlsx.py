import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
out = sys.argv[1]
BURG, IVORY, SOFT = '800020', 'F7F1E6', 'E8D5B0'
wb = Workbook(); ws = wb.active; ws.title = 'Weekly dashboard'
ws['A1'] = 'Weekly Operations Planner: dashboard'; ws['A1'].font = Font(size=16, bold=True, color=BURG)
ws['A2'] = 'Type in the shaded cells for each week. Rates calculate. Count money received, not promises. Example values are an illustration, not a target or earnings claim.'
heads = ['Measure', 'Week 1', 'Week 2', 'Week 3', 'Week 4', 'Month']
for j, h in enumerate(heads, 1):
    c = ws.cell(4, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=BURG)
inputs = [('Qualified inquiries', [8, 0, 0, 0]), ('Quotes sent', [5, 0, 0, 0]), ('Quotes accepted', [3, 0, 0, 0]), ('Invoiced ($)', [950, 0, 0, 0]), ('Collected ($)', [780, 0, 0, 0]),
          ('Jobs that passed your quality check', [3, 0, 0, 0]), ('Hours worked', [10, 0, 0, 0]), ('Supplies cost ($)', [45, 0, 0, 0]), ('Travel cost ($)', [30, 0, 0, 0])]
fill = PatternFill('solid', fgColor=IVORY); thin = Side(style='thin', color=SOFT)
for i, (lab, vals) in enumerate(inputs, start=5):
    ws.cell(i, 1, lab).font = Font(bold=True)
    for j, v in enumerate(vals, start=2):
        c = ws.cell(i, j, v); c.fill = fill; c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.cell(i, 6, f'=SUM(B{i}:E{i})')
rates = [('Quote win rate (%)', 'IF({c}6>0,{c}7/{c}6*100,"")'), ('Collection rate (%)', 'IF({c}8>0,{c}9/{c}8*100,"")'), ('Average accepted job ($)', 'IF({c}7>0,{c}8/{c}7,"")'),
         ('Outstanding ($)', '{c}8-{c}9'), ('Cost per week: supplies + travel ($)', '{c}12+{c}13'), ('Cash after supplies and travel ($)', '{c}9-({c}12+{c}13)'),
         ('Collected per hour worked ($)', 'IF({c}11>0,{c}9/{c}11,"")')]
for i, (lab, f) in enumerate(rates, start=15):
    ws.cell(i, 1, lab).font = Font(bold=True)
    for j, col in enumerate('BCDEF', start=2): ws.cell(i, j, '=' + f.format(c=col))
for r in (8, 9, 12, 13, 17, 18, 19, 20, 21): 
    for col in range(2, 7): ws.cell(r, col).number_format = '$#,##0.00'
for r in (15, 16):
    for col in range(2, 7): ws.cell(r, col).number_format = '0.0'
ws.column_dimensions['A'].width = 42
for col in 'BCDEF': ws.column_dimensions[col].width = 14
ws['A23'] = 'Collected per hour is before your own overhead and taxes. Not legal, tax or professional advice.'
pl = wb.create_sheet('Pipeline')
for j, h in enumerate(['Quote #', 'Customer', 'Service', 'Price ($)', 'Status', 'Sent', 'Decide by', 'Permission to follow up?'], 1):
    c = pl.cell(1, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=BURG); pl.column_dimensions[c.column_letter].width = 18
dv = DataValidation(type='list', formula1='"Draft,Sent,Accepted,Declined,Expired"', allow_blank=True); pl.add_data_validation(dv); dv.add('E2:E101')
dv2 = DataValidation(type='list', formula1='"Yes,No,Unknown"', allow_blank=True); pl.add_data_validation(dv2); dv2.add('H2:H101')
for r in range(2, 102): pl.cell(r, 4).number_format = '$#,##0.00'
pl['J1'] = 'Status'; pl['K1'] = 'Count'; pl['L1'] = 'Value ($)'
for i, st in enumerate(['Draft', 'Sent', 'Accepted', 'Declined', 'Expired'], start=2):
    pl.cell(i, 10, st); pl.cell(i, 11, f'=COUNTIF($E$2:$E$101,J{i})'); pl.cell(i, 12, f'=SUMIF($E$2:$E$101,J{i},$D$2:$D$101)'); pl.cell(i, 12).number_format = '$#,##0.00'
pl['J8'] = 'Open (Sent) value'; pl['L8'] = '=L3'; pl['L8'].number_format = '$#,##0.00'
inv = wb.create_sheet('Invoices')
for j, h in enumerate(['Invoice #', 'Customer', 'Amount ($)', 'Sent', 'Due', 'Paid on', 'Days past due', 'Open amount ($)'], 1):
    c = inv.cell(1, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=BURG); inv.column_dimensions[c.column_letter].width = 16
for r in range(2, 102):
    inv.cell(r, 7, f'=IF(AND(E{r}<>"",F{r}="",TODAY()>E{r}),TODAY()-E{r},"")')
    inv.cell(r, 8, f'=IF(AND(C{r}<>"",F{r}=""),C{r},0)')
    for col in (3, 8): inv.cell(r, col).number_format = '$#,##0.00'
    for col in (4, 5, 6): inv.cell(r, col).number_format = 'yyyy-mm-dd'
inv['J1'] = 'Total open ($)'; inv['K1'] = '=SUM(H2:H101)'; inv['K1'].number_format = '$#,##0.00'
inv['J2'] = 'Oldest unpaid (days)'; inv['K2'] = '=IF(COUNT(G2:G101)>0,MAX(G2:G101),0)'
wb.save(out)
