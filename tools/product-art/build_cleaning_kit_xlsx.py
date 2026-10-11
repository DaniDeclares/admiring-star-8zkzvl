import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
out = sys.argv[1]
wb = Workbook(); ws = wb.active; ws.title = 'Pricing'
BURG, IVORY, SOFT = '800020', 'F7F1E6', 'E8D5B0'
ws['A1'] = 'Cleaning Operator Job Kit: Pricing and profit check'; ws['A1'].font = Font(name='Calibri', size=16, bold=True, color=BURG)
ws['A2'] = 'Type in the shaded cells. Other cells calculate. Example values are an illustration, not a recommended price or rate.'
rows = [('Estimated labor hours', 3.5), ('Hourly labor cost', 22), ('Supplies and consumables', 9), ('Travel miles (round trip)', 14), ('Mileage rate', 0.67), ('Other costs', 0), ('Quoted price to customer', 140), ('Margin target (%)', 30)]
inp = PatternFill('solid', fgColor=IVORY); thin = Side(style='thin', color=SOFT)
for i, (l, v) in enumerate(rows, start=4):
    ws.cell(i, 1, l).font = Font(bold=True); c = ws.cell(i, 2, v); c.fill = inp; c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
calc = [('Labor cost', '=B4*B5'), ('Travel cost', '=B7*B8'), ('Total cost', '=B13+B14+B6+B9'), ('Profit', '=B10-B15'), ('Margin (%)', '=IF(B10>0,B16/B10*100,"enter a price")'), ('Decision', '=IF(B10<=0,"enter a price",IF(B17>=B11,"Meets target: go ahead","Below target: re-scope or re-price"))'),
        ('Minimum price for your target', '=IF(B11<100,B15/(1-B11/100),"target must be under 100")')]
for i, (l, f) in enumerate(calc, start=13):
    ws.cell(i, 1, l).font = Font(bold=True); ws.cell(i, 2, f)
for r in (13, 14, 15, 16, 10, 6, 9, 5, 8, 19): ws.cell(r, 2).number_format = '$#,##0.00'
ws['B17'].number_format = '0.0'
ws.column_dimensions['A'].width = 36; ws.column_dimensions['B'].width = 34
dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', showErrorMessage=True, errorTitle='Check the number', error='Enter a number of 0 or more.')
ws.add_data_validation(dv)
for r in (4, 5, 6, 7, 8, 9, 10, 11): dv.add(ws.cell(r, 2))
ws['A21'] = 'Not legal, tax or professional advice. Original planning template, licensed for one purchasing business or household.'
log = wb.create_sheet('Job log')
hdr = ['Job ID', 'Date', 'Estimated hours', 'Actual hours', 'Quoted price', 'Actual supplies', 'Actual travel', 'Labor rate', 'Profit', 'Hours variance', 'Notes']
for j, h in enumerate(hdr, 1):
    c = log.cell(1, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=BURG); c.alignment = Alignment(wrap_text=True)
    log.column_dimensions[c.column_letter].width = 16
for r in range(2, 52):
    log.cell(r, 9, f'=IF(E{r}="","",E{r}-(D{r}*H{r}+F{r}+G{r}))'); log.cell(r, 10, f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})')
    for col in (5, 6, 7, 8, 9): log.cell(r, col).number_format = '$#,##0.00'
log.freeze_panes = 'A2'; log.column_dimensions['K'].width = 40
wb.save(out)
