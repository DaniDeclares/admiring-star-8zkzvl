"""Time-and-cost validation workbook for the near-threshold services. Blank by design: no sample data ships.
Rates are DEC-LABOR-001 (owner-approved 2026-09-01): Tier 1 $60, Tier 2 $75, Tier 3 $100. Floor 50% (checkout gate).
Usage: build_validation_workbook.py out.xlsx"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
BURG, CREAM, SOFT, IVORY = '800020', 'FFFDF6', 'E8D5B0', 'F7F1E6'
wb = Workbook()
hdr = Font(bold=True, color=CREAM); hfill = PatternFill('solid', fgColor=BURG); inp = PatternFill('solid', fgColor=IVORY)
thin = Side(style='thin', color=SOFT); box = Border(left=thin, right=thin, top=thin, bottom=thin)
def head(ws, row, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v); c.font = hdr; c.fill = hfill; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = box

# --- Targets ---
t = wb.active; t.title = 'Targets'
t['A1'] = 'Services under validation (prices from Production; hours are DRAFT estimates until validated)'; t['A1'].font = Font(bold=True, size=13, color=BURG)
head(t, 3, ['SKU', 'Service', 'Locked price', 'Tier', 'Rate (DEC-LABOR-001)', 'Draft hours', 'Max hours at 50% floor', 'Draft margin', 'Status'])
SK = [('DNI-04A-001', 'Virtual Assistant', 149, 2, 1.0), ('DNI-04A-007', 'Document Formatting', 139, 2, 1.0), ('DNI-04A-009', 'Research Assistant', 225, 3, 1.5),
      ('DNI-04A-010', 'Project Coordination', 275, 3, 1.5), ('DNI-04A-008', 'Meeting Support', 225, 2, 1.5)]
for i, (sku, nm, pr, tier, hrs) in enumerate(SK, 4):
    t.cell(i, 1, sku); t.cell(i, 2, nm); t.cell(i, 3, pr); t.cell(i, 4, tier)
    t.cell(i, 5, f'=IF(D{i}=1,60,IF(D{i}=2,75,100))'); t.cell(i, 6, hrs)
    t.cell(i, 7, f'=ROUND(((1-0.5-0.029)*C{i})/E{i},2)')
    t.cell(i, 8, f'=(C{i}-F{i}*E{i}-0.029*C{i})/C{i}'); t.cell(i, 8).number_format = '0.0%'
    t.cell(i, 9, 'ECONOMICALLY_UNVERIFIED')
    for c in range(1, 10): t.cell(i, c).border = box
t['A10'] = 'Floor and fee assumptions'; t['A10'].font = Font(bold=True, color=BURG)
t['A11'] = 'Margin floor'; t['B11'] = 0.5; t['B11'].number_format = '0%'
t['A12'] = 'Processing fee used in the draft column'; t['B12'] = 0.029; t['B12'].number_format = '0.0%'
t['A13'] = 'In the Job log, enter the ACTUAL processor fee from the payout record, not this assumption.'
for col, w in zip('ABCDEFGHI', (16, 26, 13, 7, 18, 12, 22, 14, 28)): t.column_dimensions[col].width = w

# --- Job log ---
j = wb.create_sheet('Job log')
cols = ['Job date', 'SKU', 'Job reference (no names or private data)', 'Collected (USD)', 'Actual processor fee (USD)', 'Hours: intake and scoping', 'Hours: delivery', 'Hours: QA and revisions', 'Hours: admin and follow-up',
        'Other direct cost (USD)', 'Acquisition cost (USD)', 'Hours source (timer / calendar / estimate)', 'Fee source (processor record / estimate)',
        'Total hours', 'Rate', 'Labor cost', 'Total cost', 'Contribution', 'Margin', 'Clears 50%?', 'Evidence']
head(j, 1, cols)
for r in range(2, 202):
    j.cell(r, 14, f'=IF(B{r}="","",SUM(F{r}:I{r}))')
    j.cell(r, 15, f'=IF(B{r}="","",IFERROR(INDEX(Targets!$E$4:$E$8,MATCH(B{r},Targets!$A$4:$A$8,0)),"SKU?"))')
    j.cell(r, 16, f'=IF(B{r}="","",IFERROR(N{r}*O{r},""))')
    j.cell(r, 17, f'=IF(B{r}="","",IFERROR(P{r}+E{r}+J{r}+K{r},""))')
    j.cell(r, 18, f'=IF(B{r}="","",IFERROR(D{r}-Q{r},""))')
    j.cell(r, 19, f'=IF(B{r}="","",IFERROR(IF(D{r}>0,R{r}/D{r},""),""))'); j.cell(r, 19).number_format = '0.0%'
    j.cell(r, 20, f'=IF(S{r}="","",IF(S{r}>=Targets!$B$11,"YES","NO"))')
    j.cell(r, 21, f'=IF(B{r}="","",IF(AND(L{r}="timer",M{r}="processor record"),"MEASURED","PARTIAL"))')
    for c in range(1, 14): j.cell(r, c).fill = inp
    for c in range(1, 22): j.cell(r, c).border = box
dv = DataValidation(type='list', formula1='"timer,calendar,estimate"', allow_blank=True); j.add_data_validation(dv); dv.add('L2:L201')
dv2 = DataValidation(type='list', formula1='"processor record,estimate"', allow_blank=True); j.add_data_validation(dv2); dv2.add('M2:M201')
dv3 = DataValidation(type='list', formula1='=Targets!$A$4:$A$8', allow_blank=True); j.add_data_validation(dv3); dv3.add('B2:B201')
j.conditional_formatting.add('T2:T201', CellIsRule(operator='equal', formula=['"NO"'], fill=PatternFill('solid', bgColor='F4C7C3')))
j.freeze_panes = 'C2'
for i, w in enumerate((11, 14, 26, 12, 12, 11, 11, 11, 11, 12, 12, 16, 16, 9, 8, 10, 10, 11, 9, 10, 11), 1): j.column_dimensions[chr(64 + i) if i <= 26 else 'A'].width = w
j.row_dimensions[1].height = 62

# --- Summary ---
s = wb.create_sheet('Summary')
s['A1'] = 'Per-service result (only MEASURED jobs count toward verification)'; s['A1'].font = Font(bold=True, size=13, color=BURG)
head(s, 3, ['SKU', 'Service', 'Measured jobs', 'Average actual hours', 'Average margin (measured)', 'Worst margin', 'Draft hours', 'Verdict'])
for i in range(4, 9):
    k = i
    s.cell(i, 1, f'=Targets!A{k}'); s.cell(i, 2, f'=Targets!B{k}')
    s.cell(i, 3, f'=COUNTIFS(\'Job log\'!$B$2:$B$201,A{i},\'Job log\'!$U$2:$U$201,"MEASURED")')
    s.cell(i, 4, f'=IF(C{i}=0,"",AVERAGEIFS(\'Job log\'!$N$2:$N$201,\'Job log\'!$B$2:$B$201,A{i},\'Job log\'!$U$2:$U$201,"MEASURED"))'); s.cell(i, 4).number_format = '0.00'
    s.cell(i, 5, f'=IF(C{i}=0,"",AVERAGEIFS(\'Job log\'!$S$2:$S$201,\'Job log\'!$B$2:$B$201,A{i},\'Job log\'!$U$2:$U$201,"MEASURED"))'); s.cell(i, 5).number_format = '0.0%'
    s.cell(i, 6, f'=IF(C{i}=0,"",_xlfn.MINIFS(\'Job log\'!$S$2:$S$201,\'Job log\'!$B$2:$B$201,A{i},\'Job log\'!$U$2:$U$201,"MEASURED"))'); s.cell(i, 6).number_format = '0.0%'
    s.cell(i, 7, f'=Targets!F{k}')
    s.cell(i, 8, f'=IF(C{i}<3,"NEEDS MORE MEASURED JOBS (3 minimum, proposed)",IF(AND(E{i}>=Targets!$B$11,F{i}>=Targets!$B$11-0.1),"CANDIDATE FOR OWNER VERIFICATION","BELOW FLOOR: PRICE, SCOPE OR FULFILLMENT REVIEW"))')
    for c in range(1, 9): s.cell(i, c).border = box
s['A11'] = 'Rules (proposed, owner to approve): a service is not economically verified until at least 3 MEASURED jobs exist, average margin is at or above the 50% floor and no measured job is more than 10 points below it. A verdict here never changes a price, a gate, or LIVE_READY status; the owner records any verification decision.'
s['A11'].alignment = Alignment(wrap_text=True); s.merge_cells('A11:H13')
for col, w in zip('ABCDEFGH', (16, 26, 14, 16, 18, 14, 12, 52)): s.column_dimensions[col].width = w

# --- Read me ---
r = wb.create_sheet('Read me', 0)
lines = ['Time-and-cost validation workbook', '',
 'Why: 108 services are released to sell by quote then invoice, but none has measured economics. This workbook records real jobs so the 50% floor is tested on facts.',
 'Labor rate: DEC-LABOR-001 (owner-approved 2026-09-01), Tier 1 $60, Tier 2 $75, Tier 3 $100 per hour. These are internal cost rates, not wages, retail prices or provider payouts; owner draw is separate.',
 'Log: one row per paid job. Use a reference, never a client name. Take hours from a timer or calendar, and the fee from the processor record.',
 'Include acquisition cost (lead fees, ads) for the job. It is a real cost; the only measured job so far lost about 26% of revenue to it.',
 'A row counts as MEASURED only when hours come from a timer and the fee from a processor record. Estimates are PARTIAL and never verify anything.',
 'This file never changes Production. Any decision to reprice, rescope or mark a service verified belongs to the owner.']
for i, ln in enumerate(lines, 1):
    r.cell(i, 1, ln).alignment = Alignment(wrap_text=True)
r['A1'].font = Font(bold=True, size=14, color=BURG); r.column_dimensions['A'].width = 120
wb.save(sys.argv[1]); print('saved')
